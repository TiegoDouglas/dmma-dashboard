#!/usr/bin/env python3
"""Synchronize Project and Position work items from vagas.xlsx."""

from __future__ import annotations

import argparse
import json
import os
import sys
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any
from urllib.parse import quote

import requests
from openpyxl import load_workbook
from requests.auth import HTTPBasicAuth

DEFAULT_ORGANIZATION = "AccountMSFT"
DEFAULT_PROJECT = "Esteira de Vagas"
API_VERSION = "7.1"
REQUEST_TIMEOUT_SECONDS = 30
SHEET_NAME = "Planilha1"
SUPPORTED_ITEM_TYPES = ("Project", "Position")
NATURAL_KEY_FIELDS = {
    "Project": "OppID",
    "Position": "Id MyScheduling",
}
TITLE_SOURCE_FIELDS = {
    "Project": "PROJETO",
    "Position": "RoleTitle",
}
DEFAULT_FIELD_REFERENCE_NAMES = {
    "Project.OppID": "Custom.OppID_",
    "Project.Account.": "Custom.Account",
    "Project.PROJETO": "System.Title",
    "Project.Win Probability": "Custom.WinProb",
    "Project.Estimated Project Start Date": "Microsoft.VSTS.Scheduling.StartDate",
    "Project.Estimated Project End Date": "Microsoft.VSTS.Scheduling.FinishDate",
    "Position.RoleTitle": "System.Title",
}
PARENT_ITEM_TYPE = "Project"
CHILD_ITEM_TYPE = "Position"
PARENT_RELATION_TYPE = "System.LinkTypes.Hierarchy-Reverse"


class SyncError(RuntimeError):
    """Raised when workbook validation or Azure DevOps synchronization fails."""


CellValue = str | int | float | bool | date | datetime


@dataclass(frozen=True)
class ItemRecord:
    item_type: str
    source_row: int
    fields: dict[str, CellValue]

    @property
    def natural_key_name(self) -> str:
        return NATURAL_KEY_FIELDS[self.item_type]

    @property
    def natural_key_value(self) -> CellValue:
        return self.fields[self.natural_key_name]


@dataclass(frozen=True)
class RecordPair:
    project: ItemRecord
    position: ItemRecord


@dataclass(frozen=True)
class WorkbookPlan:
    pairs: tuple[RecordPair, ...]

    @property
    def projects(self) -> tuple[ItemRecord, ...]:
        return _unique_records(pair.project for pair in self.pairs)

    @property
    def positions(self) -> tuple[ItemRecord, ...]:
        return _unique_records(pair.position for pair in self.pairs)


@dataclass(frozen=True)
class FieldDefinition:
    name: str
    reference_name: str
    field_type: str
    read_only: bool


@dataclass(frozen=True)
class ResolvedRecord:
    source: ItemRecord
    fields: dict[str, CellValue]
    natural_key_reference_name: str


@dataclass(frozen=True)
class ExistingWorkItem:
    work_item_id: int
    fields: dict[str, Any]
    relations: tuple[dict[str, Any], ...]


def _unique_records(records: Any) -> tuple[ItemRecord, ...]:
    unique: dict[tuple[str, str], ItemRecord] = {}
    for record in records:
        key = (record.item_type, str(record.natural_key_value))
        previous = unique.get(key)
        if previous is not None and previous.fields != record.fields:
            raise SyncError(
                f"As linhas {previous.source_row} e {record.source_row} possuem valores "
                f"conflitantes para {record.item_type} com "
                f"{record.natural_key_name}={record.natural_key_value!r}."
            )
        unique[key] = record
    return tuple(unique.values())


def _cell_value(value: Any, field_name: str, row_number: int) -> CellValue | None:
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        return stripped or None
    if isinstance(value, (bool, int, float, datetime, date)):
        return value
    raise SyncError(
        f"O campo {field_name!r} na linha {row_number} possui tipo não suportado: "
        f"{type(value).__name__}."
    )


def read_workbook_plan(workbook_path: Path) -> WorkbookPlan:
    if not workbook_path.is_file():
        raise SyncError(f"Planilha não encontrada: {workbook_path}")

    try:
        workbook = load_workbook(
            workbook_path,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as exc:
        raise SyncError(f"Não foi possível abrir a planilha XLSX: {exc}") from exc

    try:
        if workbook.sheetnames != [SHEET_NAME]:
            raise SyncError(
                f"A planilha deve conter somente a aba {SHEET_NAME!r}; "
                f"encontrado: {', '.join(workbook.sheetnames)}."
            )

        worksheet = workbook[SHEET_NAME]
        if worksheet.max_column < 3:
            raise SyncError("Planilha1 não contém colunas suficientes para os dois tipos.")
        if worksheet.cell(1, 1).value != "Tipo de Item":
            raise SyncError("A célula A1 deve conter 'Tipo de Item'.")
        if worksheet.cell(2, 1).value != "Campos":
            raise SyncError("A célula A2 deve conter 'Campos'.")

        columns: list[tuple[int, str, str]] = []
        seen_fields: set[tuple[str, str]] = set()
        for column in range(2, worksheet.max_column + 1):
            item_type = _header_text(worksheet.cell(1, column).value, 1, column)
            field_name = _header_text(worksheet.cell(2, column).value, 2, column)
            if item_type not in SUPPORTED_ITEM_TYPES:
                raise SyncError(
                    f"Tipo de item não suportado na coluna {column}: {item_type!r}. "
                    f"Tipos aceitos: {', '.join(SUPPORTED_ITEM_TYPES)}."
                )
            field_key = (item_type, field_name.casefold())
            if field_key in seen_fields:
                raise SyncError(
                    f"O campo {field_name!r} está duplicado para {item_type}."
                )
            seen_fields.add(field_key)
            columns.append((column, item_type, field_name))

        for item_type, key_field in NATURAL_KEY_FIELDS.items():
            if (item_type, key_field.casefold()) not in seen_fields:
                raise SyncError(
                    f"O tipo {item_type} deve declarar a chave natural {key_field!r}."
                )

        pairs: list[RecordPair] = []
        for row_number in range(3, worksheet.max_row + 1):
            mapped_cells = [
                worksheet.cell(row_number, column) for column, _, _ in columns
            ]
            if not any(cell.value not in (None, "") for cell in mapped_cells):
                continue
            formula_cells = [cell.coordinate for cell in mapped_cells if cell.data_type == "f"]
            if formula_cells:
                raise SyncError(
                    f"Fórmulas não são permitidas nas células: {', '.join(formula_cells)}."
                )

            grouped: dict[str, dict[str, CellValue]] = {
                item_type: {} for item_type in SUPPORTED_ITEM_TYPES
            }
            for cell, (_, item_type, field_name) in zip(mapped_cells, columns):
                value = _cell_value(cell.value, field_name, row_number)
                if value is not None:
                    grouped[item_type][field_name] = value

            records: dict[str, ItemRecord] = {}
            for item_type in SUPPORTED_ITEM_TYPES:
                key_field = NATURAL_KEY_FIELDS[item_type]
                if key_field not in grouped[item_type]:
                    raise SyncError(
                        f"A linha {row_number} não informou a chave obrigatória "
                        f"{key_field!r} de {item_type}."
                    )
                records[item_type] = ItemRecord(
                    item_type=item_type,
                    source_row=row_number,
                    fields=grouped[item_type],
                )
            pairs.append(
                RecordPair(
                    project=records[PARENT_ITEM_TYPE],
                    position=records[CHILD_ITEM_TYPE],
                )
            )

        if not pairs:
            raise SyncError("Planilha1 não contém nenhuma linha de dados preenchida.")

        plan = WorkbookPlan(tuple(pairs))
        plan.projects
        plan.positions
        _validate_pair_relations(plan)
        return plan
    finally:
        workbook.close()


def _header_text(value: Any, row: int, column: int) -> str:
    if not isinstance(value, str) or not value.strip():
        raise SyncError(f"Cabeçalho vazio ou inválido na linha {row}, coluna {column}.")
    return value.strip()


def _validate_pair_relations(plan: WorkbookPlan) -> None:
    project_by_position: dict[tuple[str, str], tuple[str, str]] = {}
    source_row_by_position: dict[tuple[str, str], int] = {}
    for pair in plan.pairs:
        position_key = _record_key(pair.position)
        project_key = _record_key(pair.project)
        previous_project = project_by_position.get(position_key)
        if previous_project is not None and previous_project != project_key:
            raise SyncError(
                f"As linhas {source_row_by_position[position_key]} e "
                f"{pair.position.source_row} associam a mesma Position "
                f"{pair.position.natural_key_name}="
                f"{pair.position.natural_key_value!r} a Projects diferentes "
                f"({previous_project[1]!r} e {project_key[1]!r})."
            )
        project_by_position[position_key] = project_key
        source_row_by_position[position_key] = pair.position.source_row


def _display_value(value: CellValue) -> str:
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _json_value(value: CellValue) -> str | int | float | bool:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def print_dry_run(plan: WorkbookPlan, organization: str, project: str) -> None:
    print("Dry-run concluído: nenhuma chamada ao Azure DevOps foi feita.")
    print(f"Destino: {organization}/{project}")
    print(
        f"Itens únicos: {len(plan.projects)} Project(s), "
        f"{len(plan.positions)} Position(s)."
    )
    for record in (*plan.projects, *plan.positions):
        print(
            f"- {record.item_type} por {record.natural_key_name}="
            f"{_display_value(record.natural_key_value)!r}: "
            f"{len(record.fields)} campo(s)"
        )
    print(
        "Relações planejadas: Project pai de Position "
        f"({PARENT_RELATION_TYPE}), uma por linha de dados."
    )
    print(
        "Na execução real, nomes/referenceNames, tipos graváveis e identidades "
        "duplicadas serão validados pela API antes da primeira escrita."
    )


class AzureDevOpsClient:
    def __init__(
        self,
        organization: str,
        project: str,
        pat: str,
        field_overrides: dict[str, str] | None = None,
    ) -> None:
        if not pat:
            raise SyncError(
                "A variável de ambiente ADO_PAT é obrigatória. "
                "Configure-a somente como secret do GitHub Actions."
            )
        self.organization = organization
        self.project = project
        self.base_url = (
            f"https://dev.azure.com/{quote(organization, safe='')}/"
            f"{quote(project, safe='')}/_apis/wit"
        )
        self.session = requests.Session()
        self.session.auth = HTTPBasicAuth("", pat)
        self.session.headers.update({"Accept": "application/json"})
        self.field_overrides = field_overrides or {}

    def _request_json(
        self, method: str, url: str, action: str, **kwargs: Any
    ) -> dict[str, Any]:
        try:
            response = self.session.request(
                method,
                url,
                timeout=REQUEST_TIMEOUT_SECONDS,
                **kwargs,
            )
        except requests.RequestException as exc:
            raise SyncError(f"Falha ao {action}: {exc}") from exc
        if not response.ok:
            message = response.text.strip().replace("\n", " ")[:500]
            raise SyncError(
                f"Azure DevOps retornou HTTP {response.status_code} ao {action}: {message}"
            )
        try:
            payload = response.json()
        except requests.JSONDecodeError as exc:
            raise SyncError(f"Resposta inválida do Azure DevOps ao {action}.") from exc
        if not isinstance(payload, dict):
            raise SyncError(f"Resposta inesperada do Azure DevOps ao {action}.")
        return payload

    def get_field_definitions(self, item_type: str) -> list[FieldDefinition]:
        payload = self._request_json(
            "GET",
            f"{self.base_url}/workitemtypes/{quote(item_type, safe='')}/fields",
            f"consultar os campos de {item_type}",
            params={"api-version": API_VERSION},
        )
        values = payload.get("value")
        if not isinstance(values, list):
            raise SyncError(
                f"A resposta de campos de {item_type} não contém a lista 'value'."
            )
        definitions: list[FieldDefinition] = []
        for value in values:
            if not isinstance(value, dict):
                continue
            name = value.get("name")
            reference_name = value.get("referenceName")
            if isinstance(name, str) and isinstance(reference_name, str):
                definitions.append(
                    FieldDefinition(
                        name=name,
                        reference_name=reference_name,
                        field_type=str(value.get("type", "")),
                        read_only=bool(value.get("readOnly", False)),
                    )
                )
        if not definitions:
            raise SyncError(f"A API não retornou campos válidos para {item_type}.")
        return definitions

    def resolve_record(
        self, record: ItemRecord, definitions: list[FieldDefinition]
    ) -> ResolvedRecord:
        resolved: dict[str, CellValue] = {}
        source_to_reference: dict[str, str] = {}
        errors: list[str] = []
        for display_name, value in record.fields.items():
            try:
                definition = self._resolve_field(
                    record.item_type, display_name, definitions
                )
            except SyncError as exc:
                errors.append(str(exc))
                continue
            if definition.read_only:
                errors.append(
                    f"O campo {display_name!r} ({definition.reference_name}) de "
                    f"{record.item_type} é somente leitura."
                )
                continue
            if definition.reference_name in resolved:
                errors.append(
                    f"Mais de uma coluna de {record.item_type} resolve para "
                    f"{definition.reference_name}."
                )
                continue
            resolved[definition.reference_name] = value
            source_to_reference[display_name] = definition.reference_name

        if "System.Title" not in resolved:
            title_source = TITLE_SOURCE_FIELDS[record.item_type]
            title = record.fields.get(title_source)
            if title is None:
                errors.append(
                    f"{record.item_type} precisa de {title_source!r} para preencher "
                    "System.Title."
                )
            elif title_source in source_to_reference:
                resolved["System.Title"] = title

        if errors:
            raise SyncError(
                f"Falha ao validar os campos de {record.item_type} na linha "
                f"{record.source_row}:\n- "
                + "\n- ".join(errors)
                + "\nCampos graváveis disponíveis para "
                f"{record.item_type}: "
                f"{_format_field_definitions(_writable_fields(definitions))}."
            )

        return ResolvedRecord(
            source=record,
            fields=resolved,
            natural_key_reference_name=source_to_reference[record.natural_key_name],
        )

    def _resolve_field(
        self,
        item_type: str,
        display_name: str,
        definitions: list[FieldDefinition],
    ) -> FieldDefinition:
        override_key = f"{item_type}.{display_name}"
        override = self.field_overrides.get(override_key)
        default_reference = DEFAULT_FIELD_REFERENCE_NAMES.get(override_key)
        if override:
            configured_reference = override
            configured_canonical = _canonical_field_name(configured_reference)
            matches = [
                definition
                for definition in definitions
                if definition.reference_name.casefold()
                == configured_reference.casefold()
                or _canonical_field_name(definition.reference_name)
                == configured_canonical
            ]
        else:
            matches = []
            if default_reference:
                default_canonical = _canonical_field_name(default_reference)
                matches = [
                    definition
                    for definition in definitions
                    if definition.reference_name.casefold()
                    == default_reference.casefold()
                    or _canonical_field_name(definition.reference_name)
                    == default_canonical
                ]
            if not matches:
                display_canonical = _canonical_field_name(display_name)
                matches = [
                    definition
                    for definition in definitions
                    if display_canonical in _field_aliases(definition)
                ]
        unique = {definition.reference_name: definition for definition in matches}
        if len(unique) != 1:
            candidates = (
                list(unique.values())
                if unique
                else _nearby_field_candidates(display_name, definitions)
            )
            detail = (
                "não foi encontrado"
                if not unique
                else "é ambíguo após normalização canônica"
            )
            raise SyncError(
                f"O campo {display_name!r} de {item_type} {detail}. "
                f"Candidatos próximos: {_format_field_definitions(candidates)}. "
                "Corrija o processo no Azure DevOps ou configure "
                "AZURE_FIELD_REFERENCE_OVERRIDES."
            )
        return next(iter(unique.values()))

    def find_existing(self, record: ResolvedRecord) -> ExistingWorkItem | None:
        reference_name = record.natural_key_reference_name
        value = record.source.natural_key_value
        wiql = (
            "SELECT [System.Id] FROM WorkItems "
            f"WHERE [System.TeamProject] = '{_wiql_escape(self.project)}' "
            f"AND [System.WorkItemType] = '{_wiql_escape(record.source.item_type)}' "
            f"AND [{reference_name.replace(']', ']]')}] = {_wiql_literal(value)}"
        )
        payload = self._request_json(
            "POST",
            f"{self.base_url}/wiql",
            f"localizar {record.source.item_type} por {record.source.natural_key_name}",
            params={"api-version": API_VERSION},
            json={"query": wiql},
        )
        work_items = payload.get("workItems")
        if not isinstance(work_items, list):
            raise SyncError("A consulta WIQL não retornou a lista 'workItems'.")
        ids = [
            item.get("id")
            for item in work_items
            if isinstance(item, dict) and isinstance(item.get("id"), int)
        ]
        if len(ids) > 1:
            raise SyncError(
                f"Foram encontrados {len(ids)} itens {record.source.item_type} com "
                f"{record.source.natural_key_name}="
                f"{record.source.natural_key_value!r}; corrija a duplicidade."
            )
        return self.get_work_item(ids[0]) if ids else None

    def get_work_item(self, work_item_id: int) -> ExistingWorkItem:
        payload = self._request_json(
            "GET",
            f"{self.base_url}/workitems/{work_item_id}",
            f"consultar o work item {work_item_id}",
            params={"$expand": "Relations", "api-version": API_VERSION},
        )
        fields = payload.get("fields")
        relations = payload.get("relations", [])
        if not isinstance(fields, dict) or not isinstance(relations, list):
            raise SyncError(f"Resposta inválida ao consultar o work item {work_item_id}.")
        return ExistingWorkItem(
            work_item_id=work_item_id,
            fields=fields,
            relations=tuple(item for item in relations if isinstance(item, dict)),
        )

    def create_work_item(self, record: ResolvedRecord) -> int:
        patch = _field_patch(record.fields)
        payload = self._request_json(
            "POST",
            f"{self.base_url}/workitems/${quote(record.source.item_type, safe='')}",
            f"criar {record.source.item_type}",
            params={"api-version": API_VERSION},
            headers={"Content-Type": "application/json-patch+json"},
            json=patch,
        )
        work_item_id = payload.get("id")
        if not isinstance(work_item_id, int):
            raise SyncError(
                f"A API não retornou o ID do {record.source.item_type} criado."
            )
        return work_item_id

    def update_work_item(
        self, existing: ExistingWorkItem, record: ResolvedRecord
    ) -> bool:
        changed = {
            reference_name: value
            for reference_name, value in record.fields.items()
            if not _values_equal(existing.fields.get(reference_name), value)
        }
        if not changed:
            return False
        self._request_json(
            "PATCH",
            f"{self.base_url}/workitems/{existing.work_item_id}",
            f"atualizar o work item {existing.work_item_id}",
            params={"api-version": API_VERSION},
            headers={"Content-Type": "application/json-patch+json"},
            json=_field_patch(changed),
        )
        return True

    def ensure_parent_relation(
        self, child_id: int, parent_id: int, child_relations: tuple[dict[str, Any], ...]
    ) -> bool:
        expected_suffix = f"/workitems/{parent_id}".casefold()
        hierarchy_parents = [
            relation
            for relation in child_relations
            if str(relation.get("rel", "")).casefold()
            == PARENT_RELATION_TYPE.casefold()
        ]
        if any(
            str(relation.get("url", "")).casefold().endswith(expected_suffix)
            for relation in hierarchy_parents
        ):
            return False
        if hierarchy_parents:
            current_urls = ", ".join(str(item.get("url", "")) for item in hierarchy_parents)
            raise SyncError(
                f"A Position {child_id} já possui outro pai hierárquico: {current_urls}."
            )
        parent_url = (
            f"https://dev.azure.com/{quote(self.organization, safe='')}/"
            f"{quote(self.project, safe='')}/_apis/wit/workItems/{parent_id}"
        )
        patch = [
            {
                "op": "add",
                "path": "/relations/-",
                "value": {
                    "rel": PARENT_RELATION_TYPE,
                    "url": parent_url,
                    "attributes": {"comment": "Relação criada pela sincronização de vagas"},
                },
            }
        ]
        self._request_json(
            "PATCH",
            f"{self.base_url}/workitems/{child_id}",
            f"relacionar Project {parent_id} como pai de Position {child_id}",
            params={"api-version": API_VERSION},
            headers={"Content-Type": "application/json-patch+json"},
            json=patch,
        )
        return True


def synchronize(client: AzureDevOpsClient, plan: WorkbookPlan) -> None:
    definitions = {
        item_type: client.get_field_definitions(item_type)
        for item_type in SUPPORTED_ITEM_TYPES
    }
    resolved_projects: list[ResolvedRecord] = []
    resolved_positions: list[ResolvedRecord] = []
    validation_errors: list[str] = []
    for record, destination in (
        *((record, resolved_projects) for record in plan.projects),
        *((record, resolved_positions) for record in plan.positions),
    ):
        try:
            destination.append(
                client.resolve_record(record, definitions[record.item_type])
            )
        except SyncError as exc:
            validation_errors.append(str(exc))
    if validation_errors:
        raise SyncError(
            "O preflight encontrou campos inválidos; nenhuma busca ou escrita foi "
            "executada:\n\n" + "\n\n".join(validation_errors)
        )

    all_resolved = (*resolved_projects, *resolved_positions)
    existing_by_key = {
        _record_key(record.source): client.find_existing(record)
        for record in all_resolved
    }

    ids_by_key: dict[tuple[str, str], int] = {}
    relation_snapshots: dict[int, tuple[dict[str, Any], ...]] = {}
    for record in all_resolved:
        key = _record_key(record.source)
        existing = existing_by_key[key]
        if existing is None:
            work_item_id = client.create_work_item(record)
            ids_by_key[key] = work_item_id
            relation_snapshots[work_item_id] = ()
            print(
                f"Criado {record.source.item_type} {work_item_id} "
                f"({record.source.natural_key_name}="
                f"{record.source.natural_key_value!r})."
            )
        else:
            changed = client.update_work_item(existing, record)
            ids_by_key[key] = existing.work_item_id
            relation_snapshots[existing.work_item_id] = existing.relations
            action = "Atualizado" if changed else "Sem alterações"
            print(
                f"{action}: {record.source.item_type} {existing.work_item_id} "
                f"({record.source.natural_key_name}="
                f"{record.source.natural_key_value!r})."
            )

    seen_relations: set[tuple[int, int]] = set()
    for pair in plan.pairs:
        parent_id = ids_by_key[_record_key(pair.project)]
        child_id = ids_by_key[_record_key(pair.position)]
        relation_key = (parent_id, child_id)
        if relation_key in seen_relations:
            continue
        seen_relations.add(relation_key)
        created = client.ensure_parent_relation(
            child_id, parent_id, relation_snapshots.get(child_id, ())
        )
        if created:
            print(f"Relação criada: Project {parent_id} -> Position {child_id}.")
        else:
            print(f"Relação já existente: Project {parent_id} -> Position {child_id}.")


def _record_key(record: ItemRecord) -> tuple[str, str]:
    return record.item_type, str(record.natural_key_value)


def _field_patch(fields: dict[str, CellValue]) -> list[dict[str, Any]]:
    return [
        {
            "op": "add",
            "path": f"/fields/{_escape_json_pointer(reference_name)}",
            "value": _json_value(value),
        }
        for reference_name, value in fields.items()
    ]


def _escape_json_pointer(value: str) -> str:
    return value.replace("~", "~0").replace("/", "~1")


def _canonical_field_name(value: str) -> str:
    decomposed = unicodedata.normalize("NFKD", value)
    return "".join(
        character.casefold()
        for character in decomposed
        if character.isalnum()
    )


def _field_aliases(definition: FieldDefinition) -> set[str]:
    reference_leaf = definition.reference_name.rsplit(".", 1)[-1]
    return {
        _canonical_field_name(definition.name),
        _canonical_field_name(definition.reference_name),
        _canonical_field_name(reference_leaf),
    }


def _nearby_field_candidates(
    requested_name: str, definitions: list[FieldDefinition]
) -> list[FieldDefinition]:
    requested = _canonical_field_name(requested_name)
    if not requested:
        return []
    return [
        definition
        for definition in definitions
        if any(
            requested in alias or alias in requested
            for alias in _field_aliases(definition)
            if alias
        )
    ]


def _writable_fields(
    definitions: list[FieldDefinition],
) -> list[FieldDefinition]:
    return [definition for definition in definitions if not definition.read_only]


def _format_field_definitions(definitions: list[FieldDefinition]) -> str:
    if not definitions:
        return "nenhum"
    return "; ".join(
        f"{definition.name} ({definition.reference_name})"
        for definition in sorted(
            definitions,
            key=lambda item: (item.name.casefold(), item.reference_name.casefold()),
        )
    )


def _wiql_escape(value: str) -> str:
    return value.replace("'", "''")


def _wiql_literal(value: CellValue) -> str:
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, (int, float)):
        return str(value)
    return f"'{_wiql_escape(_display_value(value))}'"


def _values_equal(current: Any, desired: CellValue) -> bool:
    desired_json = _json_value(desired)
    if current == desired_json:
        return True
    if isinstance(desired, (date, datetime)) and isinstance(current, str):
        return current[:10] == desired.isoformat()[:10]
    if isinstance(desired_json, (int, float)) and isinstance(current, (int, float)):
        return float(current) == float(desired_json)
    return str(current).strip() == str(desired_json).strip()


def parse_field_overrides(raw_value: str) -> dict[str, str]:
    if not raw_value.strip():
        return {}
    try:
        value = json.loads(raw_value)
    except json.JSONDecodeError as exc:
        raise SyncError(
            "AZURE_FIELD_REFERENCE_OVERRIDES deve ser um objeto JSON válido."
        ) from exc
    if not isinstance(value, dict) or not all(
        isinstance(key, str) and isinstance(reference, str) and reference.strip()
        for key, reference in value.items()
    ):
        raise SyncError(
            "AZURE_FIELD_REFERENCE_OVERRIDES deve mapear nomes para referenceNames."
        )
    return {key: reference.strip() for key, reference in value.items()}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Sincroniza Projects e Positions de vagas.xlsx com Azure DevOps."
    )
    parser.add_argument(
        "--workbook",
        type=Path,
        default=Path(__file__).with_name("vagas.xlsx"),
        help="Caminho da planilha XLSX.",
    )
    parser.add_argument("--organization", default=DEFAULT_ORGANIZATION)
    parser.add_argument("--project", default=DEFAULT_PROJECT)
    parser.add_argument(
        "--work-item-id",
        type=int,
        help=argparse.SUPPRESS,
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Valida e exibe o plano sem acessar ou alterar o Azure DevOps.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        plan = read_workbook_plan(args.workbook)
        if args.dry_run:
            print_dry_run(plan, args.organization, args.project)
            return 0

        overrides = parse_field_overrides(
            os.getenv("AZURE_FIELD_REFERENCE_OVERRIDES", "")
        )
        client = AzureDevOpsClient(
            organization=args.organization,
            project=args.project,
            pat=os.getenv("ADO_PAT", ""),
            field_overrides=overrides,
        )
        synchronize(client, plan)
        print("Sincronização concluída com sucesso.")
        return 0
    except SyncError as exc:
        print(f"Erro: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
